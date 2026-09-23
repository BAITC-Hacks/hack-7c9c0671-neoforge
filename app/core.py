"""Local meeting transcription, diarization, action extraction and export."""
from __future__ import annotations

import re
import subprocess
import tempfile
from collections import Counter
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape

UNKNOWN = "Не указан / көрсетілмеген"
NAME = r"[А-ЯӘҒҚҢӨҰҮҺІЁ][а-яәғқңөұүһіё]+(?:\s+[А-ЯӘҒҚҢӨҰҮҺІЁ][а-яәғқңөұүһіё]+)?"
VERBS = (
    r"(?:подготов(?:ить|ьте|лю)|провест(?:и|ите)|разработать|согласовать|представить|"
    r"провер(?:ить|ьте)|организ(?:овать|уйте|уем)|най(?:ти|дите|ду)|ищите|направ(?:ить|ьте)|"
    r"обнов(?:ить|ите|лю)|доложить|разобраться|разберитесь|соберите|выставляйте|"
    r"зафиксируйте|пропишите|привлеките|свяжитесь|сделаю|возьмите|запрошу|"
    r"дайынд(?:ау|аңдар)|жаса(?:у|ңыздар)|өткіз(?:у|іңіздер)|тексер(?:у|іңіздер)|"
    r"жібер(?:у|іңіздер)|келіс(?:у|іңіздер)|ұйымдастыр(?:у|ыңыздар))"
)
FIRST_PERSON = re.compile(r"(?i)\b(?:подготовлю|найду|обновлю|сделаю|запрошу|организуем)\b")
IMPERATIVE = re.compile(r"(?i)\b(?:подготовьте|проведите|проверьте|организуйте|разберитесь|найдите|ищите|выставляйте|зафиксируйте|пропишите|направьте|обновите|привлеките|свяжитесь|возьмите)\b")
DEADLINE = re.compile(
    r"(?i)(?:до\s+\d{1,2}\s+[а-яәғқңөұүһё]+|к\s+\d{1,2}\s+[а-яәғқңөұүһё]+|"
    r"до конца недели|на этой неделе|на следующей неделе|текущая неделя|"
    r"за\s+\d+\s+(?:недел[юьиь]|дн[яей]|месяц[аев]?)|через\s+\d+\s+(?:недел[юьиь]|дн[яей])|"
    r"до\s+(?:пятницы|среды|понедельника|вторника|четверга)|к\s+(?:пятнице|среде)|"
    r"(?:осы|келесі)\s+аптада|\d{1,2}\s+[а-яәғқңөұүһё]+ға\s+дейін|бір\s+апта\s+ішінде)"
)
EXPLICIT = re.compile(r"(?:ответственн(?:ый|ая|ое|ые)|жауапты)\s*[:—-]?\s*(" + NAME + r")", re.I)
DIRECT = re.compile(r"(" + NAME + r")\s*,\s*(?:вы\s+)?(?=.{0,140}\b" + VERBS + r"\b)", re.I)
LABEL = re.compile(r"^\s*(" + NAME + r")\s*(?:\([^\n]*\))?\s*$")


@dataclass
class Segment:
    start: float
    end: float
    speaker: str
    text: str


@dataclass
class Action:
    task: str
    responsible: str
    deadline: str
    speaker: str
    evidence: str
    start: float | None = None
    needs_review: bool = False
    status: str = "draft"


def _tidy(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip(" .,:;—-")


def _valid_name(value: str) -> bool:
    stopwords = {"хорошо", "понял", "принято", "согласен", "логично", "отлично"}
    words = value.split()
    return value.casefold() not in stopwords and len(words) >= 2 and all(word[0].isupper() for word in words)


def _sentences(text: str) -> list[str]:
    return [
        _tidy(part)
        for sentence in re.split(r"(?<=[.!?])\s+", text)
        for part in re.split(r"(?i)(?=\b(?:первое|второе|третье|четвёртое|пятое)\s*:)", sentence)
        if _tidy(part)
    ]


def extract_actions(segments: list[Segment]) -> list[Action]:
    """Extract conservative candidates and retain the source evidence."""
    result: list[Action] = []
    previous_speaker: str | None = None
    for seg in segments:
        last_addressee: str | None = None
        for piece in _sentences(seg.text):
            if re.match(r"(?i)^(?:первое|второе|третье|четвёртое|пятое)\s*:", piece):
                last_addressee = None
            addressed = re.search(r"(" + NAME + r")\s*,", piece)
            if addressed and _valid_name(addressed.group(1)):
                last_addressee = addressed.group(1)
            verb = re.search(r"(?i)\b" + VERBS + r"\b", piece)
            if not verb:
                continue
            explicit = EXPLICIT.search(piece)
            direct = DIRECT.search(piece)
            direct_name = direct.group(1) if direct and _valid_name(direct.group(1)) else None
            speaker_commitment = seg.speaker if FIRST_PERSON.search(piece) and _valid_name(seg.speaker) else None
            conversational_fallback = previous_speaker if IMPERATIVE.search(piece) and previous_speaker and previous_speaker != seg.speaker and _valid_name(previous_speaker) else None
            responsible = explicit.group(1) if explicit else direct_name or last_addressee or speaker_commitment or conversational_fallback or UNKNOWN
            deadline = DEADLINE.search(piece)
            task_end = explicit.start() if explicit and explicit.start() > verb.start() else len(piece)
            task = _tidy(piece[verb.start():task_end])[:350]
            # Kazakh commonly places the verb at the end (SOV), so keep the object phrase too.
            if len(task) < 12 and verb.start() > len(piece) / 2:
                task = _tidy(piece.split(",", 1)[-1])[:350]
            if len(task) < 12:
                continue
            action = Action(
                task=task,
                responsible=responsible,
                deadline=_tidy(deadline.group(0)) if deadline else UNKNOWN,
                speaker=seg.speaker,
                evidence=piece,
                start=seg.start,
                needs_review=responsible == UNKNOWN or deadline is None,
            )
            signature = (action.task.casefold(), action.responsible.casefold())
            if not any((a.task.casefold(), a.responsible.casefold()) == signature for a in result):
                result.append(action)
        previous_speaker = seg.speaker
    return result


def reference_segments(path: Path) -> list[Segment]:
    """Parse a supplied written reference. This is not audio recognition."""
    text = path.read_text(encoding="utf-8").replace("\f", "\n").split("Саммари по ключевым пунктам")[0]
    segments: list[Segment] = []
    speaker = "UNKNOWN"
    lines: list[str] = []

    def flush() -> None:
        nonlocal lines
        if lines and speaker != "UNKNOWN":
            segments.append(Segment(0, 0, speaker, _tidy(" ".join(lines))))
        lines = []

    for raw in text.splitlines():
        line = raw.strip()
        match = LABEL.match(line)
        if match and len(match.group(1).split()) == 2 and not line.startswith(("Протокол", "Тема")):
            flush()
            speaker = match.group(1)
        elif line and not line.startswith(("Протокол совещания", "АО «", "Тема:", "Часть ", "Текст совещания")):
            lines.append(line)
    flush()
    return segments


def transcribe(audio: Path, model_dir: Path) -> list[Segment]:
    """Run multilingual Whisper strictly from a local CTranslate2 model."""
    try:
        from faster_whisper import WhisperModel
    except ImportError as exc:
        raise RuntimeError("Install the audio extras from requirements-audio.txt") from exc
    if not model_dir.is_dir():
        raise FileNotFoundError(f"Local ASR model missing: {model_dir}")
    model = WhisperModel(str(model_dir), device="cpu", compute_type="int8", local_files_only=True)
    chunks, _ = model.transcribe(str(audio), beam_size=5, word_timestamps=True, vad_filter=True)
    return [Segment(float(s.start), float(s.end), "UNKNOWN", s.text.strip()) for s in chunks if s.text.strip()]


def diarize(segments: list[Segment], audio: Path, model_dir: Path) -> list[Segment]:
    """Assign the speaker with greatest temporal overlap to every ASR segment."""
    try:
        import torchaudio
        from pyannote.audio import Pipeline
    except ImportError as exc:
        raise RuntimeError("Install the audio extras from requirements-audio.txt") from exc
    if not model_dir.is_dir():
        raise FileNotFoundError(f"Local diarization model missing: {model_dir}")
    with tempfile.TemporaryDirectory() as temp:
        wav = Path(temp) / "audio.wav"
        subprocess.run(["ffmpeg", "-nostdin", "-y", "-loglevel", "error", "-i", str(audio), "-ac", "1", "-ar", "16000", str(wav)], check=True)
        pipeline = Pipeline.from_pretrained(str(model_dir))
        waveform, sample_rate = torchaudio.load(str(wav))
        diarization = pipeline({"waveform": waveform, "sample_rate": sample_rate})
        annotation = getattr(diarization, "speaker_diarization", diarization)
        turns = [(turn.start, turn.end, label) for turn, _, label in annotation.itertracks(yield_label=True)]
    for seg in segments:
        overlap: Counter[str] = Counter()
        for start, end, label in turns:
            overlap[label] += max(0, min(seg.end, end) - max(seg.start, start))
        seg.speaker = overlap.most_common(1)[0][0] if overlap and max(overlap.values()) > 0 else "UNKNOWN"
    return segments


def make_report(segments: list[Segment], title: str) -> dict:
    summary: list[str] = []
    for segment in segments:
        for sentence in re.split(r"(?<=[.!?])\s+", segment.text):
            if re.search(r"(?i)\b(?:проблема|риск|показатель|задержка|причина|мәселе|қауіп|нәтиже)\b", sentence):
                candidate = _tidy(sentence)[:240]
                if candidate and candidate not in summary:
                    summary.append(candidate)
    actions = extract_actions(segments)
    return {
        "title": _tidy(title) or "Протокол совещания",
        "created_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "summary": summary[:8],
        "actions": [asdict(action) for action in actions],
        "segments": [asdict(segment) for segment in segments],
        "stats": {"speakers": len({s.speaker for s in segments if s.speaker != "UNKNOWN"}), "segments": len(segments), "actions": len(actions), "needs_review": sum(action.needs_review for action in actions)},
        "review_required": True,
    }


def _action_state(action: dict) -> str:
    labels = {"draft": "Черновик", "in_progress": "В работе", "done": "Выполнено"}
    value = labels.get(action.get("status", "draft"), "Черновик")
    return f"{value} · проверить" if action.get("needs_review", True) else value


def export_docx(report: dict, path: Path) -> None:
    from docx import Document
    from docx.shared import Cm, Pt

    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Cm(1.7)
    section.left_margin = section.right_margin = Cm(1.8)
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(9.5)
    doc.add_heading(report["title"], 0)
    doc.add_paragraph("ЧЕРНОВИК · Требуется проверка секретарём", style="Subtitle")
    doc.add_heading("Краткое саммари", 1)
    for item in report["summary"]:
        doc.add_paragraph(item, style="List Bullet")
    if not report["summary"]:
        doc.add_paragraph("Ключевые тезисы не выделены автоматически.")
    doc.add_heading("Поручения", 1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ["Поручение", "Ответственный", "Срок", "Статус"]):
        cell.text = text
    for action in report["actions"]:
        values = [action["task"], action["responsible"], action["deadline"], _action_state(action)]
        for cell, value in zip(table.add_row().cells, values):
            cell.text = str(value)
    doc.add_heading("Транскрипт", 1)
    for segment in report["segments"]:
        doc.add_paragraph(f"[{segment['start']:.1f}-{segment['end']:.1f}] {segment['speaker']}: {segment['text']}")
    doc.save(path)


def _font_path() -> Path:
    candidates = [Path("C:/Windows/Fonts/arial.ttf"), Path("C:/Windows/Fonts/calibri.ttf"), Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"), Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf")]
    font = next((candidate for candidate in candidates if candidate.is_file()), None)
    if font is None:
        raise RuntimeError("A Unicode font (Arial, Calibri, or DejaVu Sans) is required for PDF export")
    return font


def export_pdf(report: dict, path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font_name = "MeetingUnicode"
    if font_name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(font_name, str(_font_path())))
    body = ParagraphStyle("Body", fontName=font_name, fontSize=8.5, leading=12, textColor=colors.HexColor("#263238"), spaceAfter=5)
    title = ParagraphStyle("Title", parent=body, fontSize=18, leading=23, textColor=colors.HexColor("#0F5D66"), spaceAfter=8)
    heading = ParagraphStyle("Heading", parent=body, fontSize=12, leading=16, textColor=colors.HexColor("#0F5D66"), spaceBefore=8, spaceAfter=7)

    def paragraph(value: object, style: ParagraphStyle = body) -> Paragraph:
        return Paragraph(escape(str(value)), style)

    def footer(canvas, document) -> None:
        canvas.saveState()
        canvas.setFont(font_name, 7.5)
        canvas.setFillColor(colors.HexColor("#78909C"))
        canvas.drawString(18 * mm, 11 * mm, "Meeting Assistant · черновик")
        canvas.drawRightString(192 * mm, 11 * mm, f"{document.page}")
        canvas.restoreState()

    story = [paragraph(report["title"], title), paragraph("ЧЕРНОВИК · Требуется проверка секретарём"), Spacer(1, 4 * mm), paragraph("Краткое саммари", heading)]
    story += [paragraph("• " + item) for item in report["summary"]] or [paragraph("Ключевые тезисы не выделены автоматически.")]
    story += [paragraph("Поручения", heading)]
    rows = [[paragraph(value) for value in ["Поручение", "Ответственный", "Срок", "Статус"]]]
    rows += [[paragraph(action["task"]), paragraph(action["responsible"]), paragraph(action["deadline"]), paragraph(_action_state(action))] for action in report["actions"]]
    if len(rows) == 1:
        rows.append([paragraph("Поручения не выделены"), paragraph("-"), paragraph("-"), paragraph("Нужна")])
    table = Table(rows, colWidths=[78 * mm, 43 * mm, 34 * mm, 24 * mm], repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CFD8DC")), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E0F2F1")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
    story += [table, Spacer(1, 6 * mm), paragraph("Транскрипт", heading)]
    story += [paragraph(f"[{segment['start']:.1f}-{segment['end']:.1f}] {segment['speaker']}: {segment['text']}") for segment in report["segments"]]
    document = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=16 * mm, bottomMargin=18 * mm, title=report["title"])
    document.build(story, onFirstPage=footer, onLaterPages=footer)
